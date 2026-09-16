"""The webinar Step 5 brief to a `.pptx`.

The parser's whole risk is input shape, so the two format tests are written against *real* output
rather than invented fixtures: `manual_execution/Webinar-Package_TrafficRadius_CompetitorSynthesis.md`
is what the pipeline actually produced (a Markdown table), and the block sample is the shape
`universal-webinar-prompt.md` Step 5 actually specifies. A parser that passed only hand-written
fixtures would be a parser tuned to this file.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest

from app.services.slide_deck import (
    DeckBrand,
    brand_from_design_md,
    build_pptx,
    build_zip,
    find_slide_decks,
    slug,
)

REPO_ROOT = Path(__file__).resolve().parents[3]
REAL_PACKAGE = REPO_ROOT / "manual_execution" / "Webinar-Package_TrafficRadius_CompetitorSynthesis.md"


# The prompt's own specified shape (STEP 5's code block), which no sample file currently holds.
BLOCK_BRIEF = """# Cash Flow Masterclass for Trades

## PART 3 - FULL WEBINAR SCRIPT

Words words.

### STEP 5 - SLIDE DECK BRIEF

Slide 1: Cold open
Section: Opening Hook
Visual Direction: Full-bleed dark background, no logo yet
Headline / Title on slide: "You have already tried the obvious things."
Body content on slide: (none - single line only)
Speaker note summary: Provocative claim, no welcome.

Slide 2: The real problem
Section: Problem Frame
Visual Direction: Split screen - symptom vs cause
Headline / Title on slide: It is not a content problem.
Body content on slide: Symptom list
Cause list
Speaker note summary: Reframe from symptom to cause.

### STEP 6 - REGISTRATION PAGE COPY

Slide this is not.
"""


def test_parses_the_real_markdown_table_output():
    """The format the pipeline has actually been producing.

    This is the case a spec-only parser misses: Step 5 describes key/value blocks, and the run
    emitted a table. Returning nothing here would mean the button never appeared on real work.
    """
    decks = find_slide_decks(REAL_PACKAGE.read_text(encoding="utf-8"))

    assert len(decks) == 1
    deck = decks[0]
    assert len(deck.slides) == 14
    assert deck.slides[0].number == 1
    assert deck.slides[0].section == "Opening Hook"
    # The on-slide headline wins over the brief's internal name ("Cold open").
    assert deck.slides[0].display_title == "You've already tried the obvious things."
    assert deck.slides[13].number == 14


def test_parses_the_prompt_specified_block_format():
    decks = find_slide_decks(BLOCK_BRIEF)

    assert len(decks) == 1
    slides = decks[0].slides
    assert [s.number for s in slides] == [1, 2]
    assert slides[0].display_title == "You have already tried the obvious things."
    assert slides[0].section == "Opening Hook"
    # A value continued on the next line belongs to the field above it, not to a new slide.
    assert slides[1].body == ("Symptom list", "Cause list")


def test_section_ends_at_the_next_heading_of_the_same_level():
    """Step 6's prose must not be swept into Step 5's deck."""
    slides = find_slide_decks(BLOCK_BRIEF)[0].slides
    assert all("Slide this is not" not in s.display_title for s in slides)
    assert len(slides) == 2


def test_empty_body_markers_produce_no_bullets():
    """"(none - single line only)" is an instruction, not a bullet."""
    assert find_slide_decks(BLOCK_BRIEF)[0].slides[0].body == ()


def test_visual_direction_survives_into_the_speaker_notes():
    """Nobody here can draw "split screen", so it is carried, not approximated or dropped."""
    notes = find_slide_decks(BLOCK_BRIEF)[0].slides[1].speaker_notes()
    assert "Reframe from symptom to cause." in notes
    assert "Split screen - symptom vs cause" in notes
    assert "Problem Frame" in notes


def test_deck_is_named_from_its_enclosing_heading_not_the_step():
    """"PART 3 - FULL WEBINAR SCRIPT" names a step. The webinar's own heading names the deck."""
    assert find_slide_decks(BLOCK_BRIEF)[0].title == "Cash Flow Masterclass for Trades"


def test_several_briefs_in_one_document_become_several_decks():
    """`webinar.json`'s first field is a programme - one package per topic, so Step 5 repeats."""
    document = "# Webinar One\n\n## SLIDE DECK BRIEF\n\nSlide 1: A\nHeadline: First\n\n" "# Webinar Two\n\n## SLIDE DECK BRIEF\n\nSlide 1: B\nHeadline: Second\n"
    decks = find_slide_decks(document)

    assert [d.title for d in decks] == ["Webinar One", "Webinar Two"]
    assert decks[0].slides[0].display_title == "First"
    assert decks[1].slides[0].display_title == "Second"


def test_a_document_with_no_brief_returns_nothing_rather_than_raising():
    """A webinar run whose inputs skipped Step 5 is ordinary. The UI offers no button."""
    assert find_slide_decks("# Webinar\n\nNo brief in here at all.\n") == []


def test_a_brief_heading_with_no_parsable_slides_is_skipped():
    assert find_slide_decks("# W\n\n## SLIDE DECK BRIEF\n\nTo be written.\n") == []


# --- the file itself -------------------------------------------------------------------------


def _open(payload: bytes):
    from pptx import Presentation

    return Presentation(io.BytesIO(payload))


def test_builds_a_pptx_with_a_cover_plus_one_slide_per_brief_row():
    deck = find_slide_decks(REAL_PACKAGE.read_text(encoding="utf-8"))[0]
    presentation = _open(build_pptx(deck))

    assert len(presentation.slides) == len(deck.slides) + 1  # + cover


def test_slide_text_and_speaker_notes_actually_land_in_the_file():
    """The point of the feature: the operator stops retyping the brief into PowerPoint."""
    deck = find_slide_decks(BLOCK_BRIEF)[0]
    presentation = _open(build_pptx(deck))

    second = presentation.slides[2]  # cover, slide 1, slide 2
    text = "\n".join(s.text_frame.text for s in second.shapes if s.has_text_frame)
    assert "It is not a content problem." in text
    assert "Symptom list" in text

    assert "Reframe from symptom to cause." in second.notes_slide.notes_text_frame.text


def test_widescreen_because_a_4_3_deck_reads_as_a_mistake():
    from pptx.util import Inches

    presentation = _open(build_pptx(find_slide_decks(BLOCK_BRIEF)[0]))
    assert presentation.slide_width == Inches(13.333)


def test_brand_is_read_from_the_design_md_front_matter():
    """The colours are quoted in that front matter - a bare `#a4d36b` would open a YAML comment."""
    design_md = (
        "---\n"
        "name: Traffic Radius\n"
        "colors:\n"
        '  primary: "#a4d36b"\n'
        '  surface: "#ffffff"\n'
        '  on-surface: "#171717"\n'
        "typography:\n"
        "  h1:\n"
        '    fontFamily: "Montserrat, Helvetica, sans-serif"\n'
        "---\n"
        "# DESIGN\n"
    )
    brand = brand_from_design_md(design_md)

    assert brand.primary == "#a4d36b"
    assert brand.on_surface == "#171717"
    # A CSS stack is not a PowerPoint font name.
    assert brand.font == "Montserrat"


def test_missing_design_falls_back_to_neutral_greys_not_an_invented_brand():
    """A deck in a plausible but wrong palette is the failure `design_tokens.py` documents."""
    assert brand_from_design_md(None) == DeckBrand()
    assert brand_from_design_md("no front matter here") == DeckBrand()


@pytest.mark.parametrize("value", ["#fff", "#FFFFFF", "nonsense", ""])
def test_malformed_colours_never_break_the_build(value):
    brand = DeckBrand(primary=value, surface=value, on_surface=value)
    assert len(build_pptx(find_slide_decks(BLOCK_BRIEF)[0], brand)) > 0


def test_zip_bundles_several_decks_and_keeps_names_unique():
    payload = build_zip([("a-slides.pptx", b"one"), ("a-slides.pptx", b"two")])
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert archive.namelist() == ["a-slides.pptx", "a-slides-2.pptx"]


def test_slug_is_filename_safe_and_never_empty():
    assert slug("TRAFFIC RADIUS — WEBINAR CONTENT PACKAGE") == "traffic-radius-webinar-content-package"
    assert slug("!!!") == "webinar"


# --- the routes ------------------------------------------------------------------------------
#
# Called directly rather than over HTTP: neither route touches the database unless a `run_id` is
# supplied, so a TestClient would only add an app startup (and its DATABASE_URL) to a test about
# parsing and content types.


@pytest.mark.asyncio
async def test_list_route_reports_each_deck_and_its_filename():
    from app.routers.pipeline import SlideDeckRequest, list_slide_decks

    result = await list_slide_decks(SlideDeckRequest(text=BLOCK_BRIEF))

    assert len(result.decks) == 1
    assert result.decks[0].slide_count == 2
    assert result.decks[0].filename == "cash-flow-masterclass-for-trades-slides.pptx"
    # One deck downloads as itself, not as a zip of one.
    assert result.bundle_filename == result.decks[0].filename


@pytest.mark.asyncio
async def test_a_document_with_no_brief_is_a_422_that_says_what_to_do():
    """Not a 500 and not an empty file: the operator can act on this one."""
    from fastapi import HTTPException

    from app.routers.pipeline import SlideDeckRequest, list_slide_decks

    with pytest.raises(HTTPException) as excinfo:
        await list_slide_decks(SlideDeckRequest(text="# Webinar\n\nNothing here.\n"))

    assert excinfo.value.status_code == 422
    assert "Step 5" in excinfo.value.detail


@pytest.mark.asyncio
async def test_download_returns_a_pptx_for_a_single_deck():
    from app.routers.pipeline import SlideDeckRequest, download_slide_deck

    response = await download_slide_deck(SlideDeckRequest(text=BLOCK_BRIEF))

    assert response.media_type.endswith("presentationml.presentation")
    assert "cash-flow-masterclass-for-trades-slides.pptx" in response.headers["Content-Disposition"]
    assert len(_open(response.body).slides) == 3  # cover + 2


@pytest.mark.asyncio
async def test_download_returns_a_zip_when_the_document_holds_several():
    from app.routers.pipeline import SlideDeckRequest, download_slide_deck

    document = (
        "# Webinar One\n\n## SLIDE DECK BRIEF\n\nSlide 1: A\nHeadline: First\n\n"
        "# Webinar Two\n\n## SLIDE DECK BRIEF\n\nSlide 1: B\nHeadline: Second\n"
    )
    response = await download_slide_deck(SlideDeckRequest(text=document))

    assert response.media_type == "application/zip"
    with zipfile.ZipFile(io.BytesIO(response.body)) as archive:
        assert archive.namelist() == ["webinar-one-slides.pptx", "webinar-two-slides.pptx"]


@pytest.mark.asyncio
async def test_an_index_past_the_end_is_a_422_naming_the_count():
    from fastapi import HTTPException

    from app.routers.pipeline import SlideDeckRequest, download_slide_deck

    with pytest.raises(HTTPException) as excinfo:
        await download_slide_deck(SlideDeckRequest(text=BLOCK_BRIEF), index=7)

    assert excinfo.value.status_code == 422
    assert "1 slide deck brief" in excinfo.value.detail


@pytest.mark.asyncio
async def test_a_missing_run_never_blocks_the_download():
    """A webinar is a `BRAND_THEME_STAGES` member and may legitimately have no captured design.
    Withholding the deliverable to protect a colour would be the wrong trade."""
    from app.routers.pipeline import SlideDeckRequest, download_slide_deck

    response = await download_slide_deck(SlideDeckRequest(text=BLOCK_BRIEF, run_id="not-a-uuid"))
    assert len(response.body) > 0
