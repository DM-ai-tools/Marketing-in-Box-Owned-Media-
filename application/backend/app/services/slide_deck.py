"""The webinar stage's SLIDE DECK BRIEF, turned into a real `.pptx`.

Step 5 of `universal-webinar-prompt.md` produces a slide-by-slide brief for "the designer or
presenter" — which, until now, was where it stopped. The operator got a description of a deck and
had to build the deck by hand, retyping fourteen slides and their speaker notes into PowerPoint.
This module does that transcription.

What it is not: a designer. It lays out title, body and notes on the client's palette and stops
there. The brief's `Visual Direction` line ("split-screen symptom vs cause", "4-icon row") describes
artwork nobody here can draw, so it is carried into the speaker notes verbatim rather than
approximated — a slide that *says* what it should look like is useful to the person opening the
file, and a slide that guesses is a slide they have to undo.

**Two input formats, both real.** The prompt specifies a key/value block per slide, and that is what
Step 5 asks for. What real runs actually produced is a Markdown *table* — see
`manual_execution/Webinar-Package_TrafficRadius_CompetitorSynthesis.md`, which is one wide table
with a row per slide. A parser that handled only the specified shape would return nothing on the
output the pipeline has actually been generating, so both are parsed and `tests/test_slide_deck.py`
pins each against a real sample.

**One document can hold several decks.** `webinar.json`'s first field is a *programme*: the prompt
"builds one complete webinar package per topic", so Steps 3-8 repeat per webinar and a run's
document can carry several briefs. `find_slide_decks` returns all of them, which is why the route
offers a zip when there is more than one.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

# The heading that opens the section, however it is numbered. Real outputs have called it both
# "STEP 5 — SLIDE DECK BRIEF" (the prompt's own numbering) and "PART 4 — SLIDE DECK BRIEF", so the
# phrase is matched and the numbering ignored.
_BRIEF_HEADING = re.compile(r"^(#{1,6})\s*(.*\bSLIDE\s+DECK\s+BRIEF\b.*?)\s*$", re.IGNORECASE)
_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*$")
_SLIDE_START = re.compile(r"^\**\s*Slide\s*(\d+)\s*\**\s*[:.—-]\s*(.*?)\s*$", re.IGNORECASE)
_HEX = re.compile(r"^#(?:[0-9a-fA-F]{3}|[0-9a-fA-F]{6})$")

# `Body content on slide:` and friends. The prompt's labels, plus the shortenings real output uses.
_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "section": ("section",),
    "visual": ("visual direction", "visual", "design direction", "slide cue"),
    "headline": ("headline / title on slide", "headline on slide", "headline", "title on slide"),
    "body": ("body content on slide", "body content", "body", "content"),
    "notes": ("speaker note summary", "speaker notes", "speaker note", "notes", "note"),
    "title": ("slide title", "title"),
    "number": ("slide", "slide number", "#", "no", "no."),
}

# Emitted when the brief says a slide carries no body at all — "(none — single line only)" and the
# like. Recognised rather than rendered: a bullet reading "(none)" is worse than an empty slide.
_EMPTY_MARKERS = ("none", "n/a", "na", "-", "—", "–", "")


def _clean(text: str) -> str:
    """Strip the Markdown a brief is written in, leaving the words that go on the slide."""
    out = text.strip()
    out = re.sub(r"`([^`]*)`", r"\1", out)
    out = re.sub(r"\*\*([^*]*)\*\*", r"\1", out)
    out = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", out)
    out = re.sub(r"^[-*•]\s+", "", out)
    return out.strip()


def _is_empty_value(text: str) -> bool:
    stripped = _clean(text).strip("()").strip().lower()
    if stripped in _EMPTY_MARKERS:
        return True
    # "(none — single line only)" carries an explanation after the marker.
    return bool(re.match(r"^(none|n/?a)\b", stripped))


def _canonical_field(label: str) -> str | None:
    key = _clean(label).rstrip(":").strip().lower()
    for canonical, aliases in _FIELD_ALIASES.items():
        if key in aliases:
            return canonical
    return None


@dataclass(frozen=True)
class Slide:
    number: int
    title: str
    section: str = ""
    visual: str = ""
    headline: str = ""
    body: tuple[str, ...] = ()
    notes: str = ""

    @property
    def display_title(self) -> str:
        """What goes in the title placeholder.

        The on-slide headline wins over the brief's internal slide name when both exist: `Slide 1:
        Cold open` is a label for the designer, while `"You've already tried the obvious things."`
        is the text the audience reads. Falling back the other way keeps a title on slides where the
        brief only named them.
        """
        return self.headline or self.title or f"Slide {self.number}"

    def speaker_notes(self) -> str:
        """Everything the deck knows about this slide that is not on its face.

        The visual direction goes here rather than being dropped: it is an instruction to whoever
        opens the file, and the notes pane is where a presenter looks.
        """
        parts: list[str] = []
        if self.notes:
            parts.append(self.notes)
        if self.visual:
            parts.append(f"Visual direction: {self.visual}")
        if self.section:
            parts.append(f"Section: {self.section}")
        if self.headline and self.title and self.headline != self.title:
            parts.append(f"Slide name in brief: {self.title}")
        return "\n\n".join(parts)


@dataclass
class SlideDeck:
    title: str
    slides: list[Slide] = field(default_factory=list)
    #: Prose that sat under the heading but outside the table or blocks — usually the design-token
    #: note ("reuses the client's tokens: primary #a4d36b...") and the "design rules applied" line.
    notes: list[str] = field(default_factory=list)


def _split_table_row(line: str) -> list[str]:
    row = line.strip()
    if row.startswith("|"):
        row = row[1:]
    if row.endswith("|"):
        row = row[:-1]
    return [cell.strip() for cell in row.split("|")]


def _is_divider(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c.strip() or "-") for c in cells)


def _bullets(raw: str) -> tuple[str, ...]:
    """Body copy to slide bullets.

    Split on the separators briefs actually use — real line breaks, `<br>`, bullet glyphs and
    semicolons — but *not* on commas or full stops: "9-day approval, 2-day news cycle." is one line
    on one slide, and splitting it would invent two.
    """
    if _is_empty_value(raw):
        return ()
    text = re.sub(r"<br\s*/?>", "\n", raw, flags=re.IGNORECASE)
    pieces = re.split(r"\n|(?<!^)\s+[•]\s+|;\s+", text)
    return tuple(c for c in (_clean(p) for p in pieces) if c and not _is_empty_value(c))


def _parse_table(lines: list[str]) -> list[Slide]:
    """A Markdown table whose header names the brief's fields, one row per slide."""
    rows = [ln for ln in lines if ln.strip().startswith("|")]
    if len(rows) < 2:
        return []

    header = [_canonical_field(c) for c in _split_table_row(rows[0])]
    if "title" not in header and "headline" not in header:
        return []

    slides: list[Slide] = []
    auto_number = 0
    for line in rows[1:]:
        cells = _split_table_row(line)
        if _is_divider(cells):
            continue
        values: dict[str, str] = {}
        for index, cell in enumerate(cells):
            if index >= len(header):
                break
            name = header[index]
            if name:
                values[name] = cell

        title = _clean(values.get("title", ""))
        headline = _clean(values.get("headline", ""))
        if not (title or headline):
            continue

        auto_number += 1
        raw_number = _clean(values.get("number", ""))
        number = int(raw_number) if raw_number.isdigit() else auto_number

        slides.append(
            Slide(
                number=number,
                title=title,
                section=_clean(values.get("section", "")),
                visual=_clean(values.get("visual", "")),
                headline="" if _is_empty_value(headline) else headline.strip('"“”'),
                body=_bullets(values.get("body", "")),
                notes=_clean(values.get("notes", "")),
            )
        )
    return slides


def _parse_blocks(lines: list[str]) -> list[Slide]:
    """The prompt's own shape: `Slide N: Title` followed by `Field: value` lines."""
    slides: list[Slide] = []
    current: dict[str, str] | None = None
    number = 0
    title = ""
    last_field: str | None = None

    def flush() -> None:
        nonlocal current
        if current is None:
            return
        headline = current.get("headline", "").strip('"“”')
        slides.append(
            Slide(
                number=number,
                title=title,
                section=current.get("section", ""),
                visual=current.get("visual", ""),
                headline="" if _is_empty_value(headline) else headline,
                body=_bullets(current.get("body", "")),
                notes=current.get("notes", ""),
            )
        )
        current = None

    for raw in lines:
        line = raw.rstrip()
        start = _SLIDE_START.match(line.strip())
        if start:
            flush()
            number = int(start.group(1))
            title = _clean(start.group(2))
            current = {}
            last_field = None
            continue

        if current is None:
            continue

        stripped = _clean(line)
        if not stripped:
            last_field = None
            continue

        head, sep, tail = stripped.partition(":")
        name = _canonical_field(head) if sep else None
        if name and name not in ("number", "title"):
            current[name] = tail.strip()
            last_field = name
        elif name in ("number", "title") and sep:
            title = title or tail.strip()
            last_field = None
        elif last_field:
            # A wrapped value, or a further bullet under `Body content on slide:`.
            current[last_field] = f"{current[last_field]}\n{stripped}".strip()

    flush()
    return slides


def _deck_title(lines: list[str], heading_index: int, fallback_index: int) -> str:
    """Name the deck from the nearest enclosing heading.

    A brief sits under its webinar's own heading, so the closest heading *above* it at a shallower
    level is the webinar it belongs to. That is what an operator with three decks in a folder needs
    on the filename; a numbered fallback covers a document with no such heading.
    """
    own = _BRIEF_HEADING.match(lines[heading_index])
    own_level = len(own.group(1)) if own else 6

    for index in range(heading_index - 1, -1, -1):
        match = _HEADING.match(lines[index])
        if not match:
            continue
        if len(match.group(1)) < own_level:
            name = _clean(match.group(2))
            # "PART 3 — FULL WEBINAR SCRIPT" names a step, not a webinar.
            if not re.match(r"^(part|step|output)\b", name, re.IGNORECASE):
                return name
    return f"Webinar {fallback_index}"


def find_slide_decks(document: str) -> list[SlideDeck]:
    """Every SLIDE DECK BRIEF in one webinar document, in order.

    Returns `[]` rather than raising when the document has no brief — a webinar run whose inputs
    skipped Step 5 is an ordinary outcome, and the caller offers no download instead of an error.
    """
    lines = document.splitlines()
    decks: list[SlideDeck] = []

    heading_indexes = [i for i, line in enumerate(lines) if _BRIEF_HEADING.match(line)]
    for position, start in enumerate(heading_indexes, start=1):
        level = len(_BRIEF_HEADING.match(lines[start]).group(1))

        end = len(lines)
        for index in range(start + 1, len(lines)):
            match = _HEADING.match(lines[index])
            if match and len(match.group(1)) <= level:
                end = index
                break

        body = lines[start + 1 : end]
        slides = _parse_table(body) or _parse_blocks(body)
        if not slides:
            continue

        notes = [
            _clean(ln)
            for ln in body
            if ln.strip()
            and not ln.strip().startswith("|")
            and not _SLIDE_START.match(ln.strip())
            and _canonical_field(ln.partition(":")[0]) is None
        ]
        decks.append(
            SlideDeck(
                title=_deck_title(lines, start, position),
                slides=slides,
                notes=[n for n in notes if n],
            )
        )
    return decks


@dataclass(frozen=True)
class DeckBrand:
    """Just enough of the run's DESIGN.md to stop the deck looking like stock PowerPoint."""

    primary: str = "#1F2937"
    surface: str = "#FFFFFF"
    on_surface: str = "#111827"
    font: str = "Calibri"

    @property
    def primary_rgb(self) -> RGBColor:
        return _rgb(self.primary, "1F2937")

    @property
    def surface_rgb(self) -> RGBColor:
        return _rgb(self.surface, "FFFFFF")

    @property
    def on_surface_rgb(self) -> RGBColor:
        return _rgb(self.on_surface, "111827")


def _rgb(value: str, fallback: str) -> RGBColor:
    hexcode = value.lstrip("#")
    if len(hexcode) == 3:
        hexcode = "".join(c * 2 for c in hexcode)
    if len(hexcode) != 6:
        hexcode = fallback
    try:
        return RGBColor.from_string(hexcode.upper())
    except ValueError:
        return RGBColor.from_string(fallback)


def brand_from_design_md(design_md: str | None) -> DeckBrand:
    """Read the palette out of a run's DESIGN.md front matter.

    Hand-parsed for the same reason `design_md.py` hand-writes it: nothing in this runtime depends
    on a YAML library, and the block is two levels deep. Colours are quoted in that front matter —
    a bare `#a4d36b` would open a YAML comment — so the quotes come off here.

    Defaults are deliberately neutral greys rather than an invented brand colour. A deck in the
    wrong palette is the failure `design_tokens.py` documents for generated pages: plausible, and
    wrong in a way only somebody who knows the brand will catch.
    """
    brand = DeckBrand()
    if not design_md:
        return brand

    lines = design_md.splitlines()
    if not lines or lines[0].strip() != "---":
        return brand

    colors: dict[str, str] = {}
    font = ""
    section = ""
    for line in lines[1:]:
        if line.strip() == "---":
            break
        if not line.startswith(" "):
            section = line.split(":", 1)[0].strip()
            continue
        key, sep, raw = line.strip().partition(":")
        if not sep:
            continue
        value = raw.strip().strip("\"'")
        if section == "colors" and _HEX.match(value):
            colors[key.strip()] = value
        elif not font and key.strip() == "fontFamily" and value:
            # The token carries a full CSS stack; PowerPoint wants one family name.
            font = value.split(",")[0].strip().strip("\"'")

    return DeckBrand(
        primary=colors.get("primary", brand.primary),
        surface=colors.get("surface", brand.surface),
        on_surface=colors.get("on-surface", brand.on_surface),
        font=font or brand.font,
    )


_WIDE = Inches(13.333)
_TALL = Inches(7.5)


def _set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _text_frame(slide, left: Emu, top: Emu, width: Emu, height: Emu):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    return frame


def build_pptx(deck: SlideDeck, brand: DeckBrand | None = None) -> bytes:
    """One deck to `.pptx` bytes.

    Blank layouts and explicit text boxes rather than the template's title/content placeholders:
    the stock layouts carry Office's own theme fonts and colours, and a deck that is half the
    client's palette and half Calibri-blue reads as a mistake rather than a draft.
    """
    brand = brand or DeckBrand()
    presentation = Presentation()
    presentation.slide_width = _WIDE
    presentation.slide_height = _TALL
    blank = presentation.slide_layouts[6]

    margin = Inches(0.9)
    width = _WIDE - margin * 2

    # Cover.
    cover = presentation.slides.add_slide(blank)
    _set_background(cover, brand.primary_rgb)
    frame = _text_frame(cover, margin, Inches(2.6), width, Inches(2.2))
    run = frame.paragraphs[0].add_run()
    run.text = deck.title
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.name = brand.font
    run.font.color.rgb = brand.surface_rgb

    subtitle = frame.add_paragraph()
    sub_run = subtitle.add_run()
    sub_run.text = f"{len(deck.slides)} slides — generated from the webinar slide deck brief"
    sub_run.font.size = Pt(15)
    sub_run.font.name = brand.font
    sub_run.font.color.rgb = brand.surface_rgb
    if deck.notes:
        cover.notes_slide.notes_text_frame.text = "\n\n".join(deck.notes)

    for slide_spec in deck.slides:
        slide = presentation.slides.add_slide(blank)
        _set_background(slide, brand.surface_rgb)

        # Section eyebrow, so the deck keeps the brief's structure visible while it is edited.
        if slide_spec.section:
            eyebrow = _text_frame(slide, margin, Inches(0.55), width, Inches(0.4))
            run = eyebrow.paragraphs[0].add_run()
            run.text = slide_spec.section.upper()
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.name = brand.font
            run.font.color.rgb = brand.primary_rgb

        title_frame = _text_frame(slide, margin, Inches(1.1), width, Inches(1.6))
        run = title_frame.paragraphs[0].add_run()
        run.text = slide_spec.display_title
        run.font.size = Pt(32)
        run.font.bold = True
        run.font.name = brand.font
        run.font.color.rgb = brand.on_surface_rgb

        if slide_spec.body:
            body_frame = _text_frame(slide, margin, Inches(2.9), width, Inches(3.4))
            for index, line in enumerate(slide_spec.body):
                paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
                bullet = paragraph.add_run()
                bullet.text = f"•  {line}" if len(slide_spec.body) > 1 else line
                bullet.font.size = Pt(20)
                bullet.font.name = brand.font
                bullet.font.color.rgb = brand.on_surface_rgb
                paragraph.space_after = Pt(12)

        notes = slide_spec.speaker_notes()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def slug(text: str) -> str:
    out = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return out[:60] or "webinar"


def build_zip(decks: list[tuple[str, bytes]]) -> bytes:
    """Several decks as one download. Names are made unique here, not assumed to be."""
    buffer = io.BytesIO()
    seen: dict[str, int] = {}
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in decks:
            count = seen.get(name, 0)
            seen[name] = count + 1
            final = name if count == 0 else name.replace(".pptx", f"-{count + 1}.pptx")
            archive.writestr(final, payload)
    return buffer.getvalue()
